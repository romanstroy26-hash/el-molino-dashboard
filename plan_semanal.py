"""Los dos planes SEMANALES del dueño, hechos por el sistema: PANADERÍA (por base, con panes salados) y DECORACIÓN (por sabor, sin salados).

Los PDF de ejemplo ("Plan_Panaderia_Semana_<lunes>.pdf" y "Plan_Decoracion_Semana_<lunes>.pdf" en data/planes_produccion/<punto>) dan SOLO
el formato: título, filas y grupos en su orden, etiquetas de las opciones (38/42/46) y el texto del pie. Ninguna cantidad de ellos se usa:
todas salen de los planes diarios del sistema (plan_propio.plan_del_dia), que se hacen solo con ventas.

  - Mat = lo que se hornea de 7 a 12; Vesp = de 13 a 17 (suma de las columnas de horneado del plan diario).
  - OPCIONES 1/2/3 = unidades que se añaden después de las 5pm según las ventas a la 1pm (iguales todos los días, como en la hoja).
  - DECORACIÓN = una fila por sabor (producto del catálogo); PANADERÍA = una fila por base: la suma de los sabores que le tocan, más los
    panes salados. Por eso "mismas cantidades, separadas por sabor": ambos planes salen de las mismas cifras.
  - RESERVA REFRI (diaria) es una decisión manual: se escribe a mano, el sistema no la calcula.

Puro salvo la lectura de los PDF de formato (pdfplumber) y las salidas (pymupdf / openpyxl).
"""

from __future__ import annotations

import datetime as dt
import difflib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import plan_dia
import plan_propio as pp

DIAS_CORTOS = ("LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO")
UMBRAL_PARECIDO = 0.85
# nombre de caja/hoja que no se reconoce por palabras: el "Pan de Muerto Negro" de la hoja de panadería es el de chocolate (supuesto del
# sistema: la masa negra es la de chocolate; si no es así, cambiar aquí)
BASE_ESPECIAL = {plan_dia.clave_producto("Pan de Muerto Chocolate"): plan_dia.clave_producto("Pan de Muerto Negro")}


# --------------------------------------------------------------------------------------------- formato (lectura de los PDF)

@dataclass
class Formato:
    tipo: str                    # "panaderia" | "decoracion"
    titulo: str                  # "PLAN SEMANAL PANADERÍA — BASES"
    columna: str                 # "BASE" | "SABOR"
    etiquetas_opciones: tuple    # ("1\n38", "2\n42", "3\n46")
    titulo_opciones: str         # "OPCIONES (1pm)"
    filas: list                  # [("grupo", "MINI CROISSANT"), ("item", "Mini Croissant Frutos"), ...]
    pie: str
    reserva: str | None = None   # "RESERVA\nREFRI\n(diaria)" si la hoja trae esa columna
    fuente: str = ""

    @property
    def umbrales(self) -> tuple:
        return pp.umbrales_de_etiquetas(self.etiquetas_opciones)

    @property
    def items(self) -> list:
        return [n for t, n in self.filas if t == "item"]


def leer_formato(path: Path) -> Formato | None:
    """Estructura de una hoja semanal (sin cantidades). None si no tiene el formato esperado."""
    import pdfplumber

    path = Path(path)
    titulo = columna = ""
    opciones: tuple = ()
    titulo_opciones = "OPCIONES (1pm)"
    reserva = None
    pie = ""
    filas: list = []
    tipo = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for tabla in page.extract_tables():
                for r in tabla:
                    c = [(x or "").strip() for x in r]
                    if not c or not c[0] and not any(c):
                        continue
                    primero = c[0]
                    resto_vacio = not any(c[1:])
                    if "PLAN SEMANAL" in primero.upper():
                        if not titulo:
                            titulo = re.split(r"\s+—\s+semana", primero)[0].strip()
                            tipo = "panaderia" if "PANADER" in primero.upper() else "decoracion"
                        continue
                    if primero.upper() in ("BASE", "SABOR"):
                        if not columna:
                            columna = primero.upper()
                            titulo_opciones = next((x for x in c if x.upper().startswith("OPCIONES")), titulo_opciones)
                            if len(c) >= 19 and c[18]:
                                reserva = c[18]
                        continue
                    if primero == "" and len(c) > 17 and c[1].lower() == "mat":
                        if not opciones:
                            opciones = tuple(c[15:18])
                        continue
                    if primero.upper() == "TOTAL":
                        continue
                    if len(primero) > 60 and resto_vacio:
                        pie = primero
                        continue
                    if primero and resto_vacio and primero.upper() == primero:
                        filas.append(("grupo", primero))
                    elif primero:
                        filas.append(("item", primero))
    if not titulo or not filas or len(opciones) != 3:
        return _leer_formato_texto(path)                  # PDF sin celdas bien marcadas (p. ej. los que escribe este mismo sistema)
    return Formato(tipo=tipo, titulo=titulo, columna=columna or ("BASE" if tipo == "panaderia" else "SABOR"), etiquetas_opciones=opciones,
                   titulo_opciones=titulo_opciones, filas=filas, pie=_sin_nota_del_sistema(pie), reserva=reserva, fuente=path.name)


def _sin_nota_del_sistema(pie: str) -> str:
    """Quita la línea "Plan hecho por el sistema..." que este módulo agrega al pie, para que no se acumule si el PDF se usa de formato."""
    return re.split(r"\s*·?\s*Plan hecho por el sistema", pie)[0].strip()


def _leer_formato_texto(path: Path) -> Formato | None:
    """Lo mismo que leer_formato pero por líneas de texto (grupos en mayúsculas sin cifras, productos con sus cifras detrás)."""
    import pdfplumber

    titulo = columna = ""
    tipo = ""
    opciones: tuple = ()
    reserva = None
    filas: list = []
    pie: list = []
    with pdfplumber.open(Path(path)) as pdf:
        for page in pdf.pages:
            lineas = [x.strip() for x in (page.extract_text() or "").splitlines() if x.strip()]
            en_pie = False
            for linea in lineas:
                if "PLAN SEMANAL" in linea.upper():
                    if not titulo:
                        titulo = re.split(r"\s+—\s+semana", linea)[0].strip()
                        tipo = "panaderia" if "PANADER" in linea.upper() else "decoracion"
                    continue
                if re.match(r"^(BASE|SABOR)\b", linea):
                    if not columna:
                        columna = linea.split()[0]
                        m = re.search(r"\b1\s+(\d+(?:[.,]\d+)?)\s+2\s+(\d+(?:[.,]\d+)?)\s+3\s+(\d+(?:[.,]\d+)?)", linea)
                        if m:
                            opciones = tuple(f"{i + 1}\n{m.group(i + 1)}" for i in range(3))
                        if "RESERVA" in linea.upper():
                            reserva = "RESERVA\nREFRI\n(diaria)"
                    continue
                if linea.startswith("LUNES") or linea.startswith("Mat Vesp"):
                    continue
                if linea.upper().startswith("TOTAL"):
                    en_pie = True
                    continue
                if en_pie:
                    pie.append(linea)
                    continue
                nombre = re.sub(r"(\s+\d+)+$", "", linea).strip()
                if nombre and nombre.upper() == nombre and not any(c.isdigit() for c in nombre):
                    filas.append(("grupo", nombre))
                elif nombre:
                    filas.append(("item", nombre))
    if not titulo or not filas or len(opciones) != 3:
        return None
    return Formato(tipo=tipo, titulo=titulo, columna=columna or ("BASE" if tipo == "panaderia" else "SABOR"), etiquetas_opciones=opciones,
                   titulo_opciones="OPCIONES (1pm)", filas=filas, pie=_sin_nota_del_sistema(" ".join(pie)), reserva=reserva, fuente=Path(path).name)


def formatos_en_carpeta(carpeta: Path, _cache: dict = {}) -> dict:
    """{"panaderia": Formato, "decoracion": Formato}: la hoja MÁS NUEVA de cada tipo en la carpeta (por fecha de modificación)."""
    carpeta = Path(carpeta)
    salida: dict = {}
    if not carpeta.exists():
        return salida
    for patron in ("Plan_Panaderia_Semana_*.pdf", "Plan_Decoracion_Semana_*.pdf"):
        archivos = sorted(carpeta.glob(patron), key=lambda p: p.stat().st_mtime)
        for p in reversed(archivos):
            clave = (str(p), p.stat().st_mtime)
            if clave not in _cache:
                try:
                    _cache[clave] = leer_formato(p)
                except Exception:  # noqa: BLE001 -- un PDF dañado no debe tumbar la página
                    _cache[clave] = None
            f = _cache[clave]
            if f is not None:
                salida[f.tipo] = f
                break
    return salida


# --------------------------------------------------------------------------------------------- mapa catálogo <-> filas de la hoja

def _clave_fila(nombre: str) -> str:
    return plan_dia.clave_producto(re.sub(r"\(.*?\)", " ", nombre))


def mapa_decoracion(formato: Formato, claves_catalogo: list) -> dict:
    """{nombre de la fila: clave del catálogo | None}: igual de nombre o casi igual (erratas de la hoja)."""
    salida = {}
    for nombre in formato.items:
        k = _clave_fila(nombre)
        if k in claves_catalogo:
            salida[nombre] = k
            continue
        cerca = difflib.get_close_matches(k, claves_catalogo, n=1, cutoff=UMBRAL_PARECIDO)
        salida[nombre] = cerca[0] if cerca else None
    return salida


def mapa_panaderia(formato: Formato, items: list) -> dict:
    """{nombre de la base: [claves del catálogo]}: a cada producto le toca la base más específica cuyas palabras están todas en el suyo
    ("Croissant Fresa" -> Croissant, "Mini Croissant Fresa" -> Mini Croissant); si hay empate gana la base que contiene la primera palabra
    de su nombre ("Pizza Salchicha Tocino" -> Pizza, no Salchicha); un producto sin base queda fuera. `items` = productos del catálogo."""
    bases = {nombre: frozenset(_clave_fila(nombre).split()) for nombre in formato.items}
    clave_de_base = {_clave_fila(n): n for n in formato.items}
    salida: dict = {n: [] for n in formato.items}
    for item in items:
        k = item.clave
        especial = clave_de_base.get(BASE_ESPECIAL.get(k, ""))
        if especial:
            salida[especial].append(k)
            continue
        palabras = frozenset(k.split())
        candidatas = [(len(p), n) for n, p in bases.items() if p and p <= palabras]
        if candidatas:
            tope = max(c[0] for c in candidatas)
            mejores = [n for t, n in candidatas if t == tope]
            primera = (plan_dia._norm(item.producto).split() or [""])[0]
            elegida = next((n for n in mejores if primera in bases[n]), mejores[0])
            salida[elegida].append(k)
            continue
        cerca = difflib.get_close_matches(k, [_clave_fila(n) for n in formato.items], n=1, cutoff=UMBRAL_PARECIDO)
        if cerca:
            salida[clave_de_base[cerca[0]]].append(k)
    return salida


# --------------------------------------------------------------------------------------------- construcción

@dataclass
class FilaSemana:
    nombre: str
    grupo: bool = False
    mat: list = field(default_factory=lambda: [None] * 7)
    vesp: list = field(default_factory=lambda: [None] * 7)
    opciones: tuple = (None, None, None)
    reserva: int | None = None


@dataclass
class PlanSemana:
    formato: Formato
    lunes: dt.date
    filas: list
    colchon: float
    faltan_dias: list = field(default_factory=list)       # días de la semana sin plan (falta historia)
    sin_producto: list = field(default_factory=list)      # filas de la hoja sin producto en el catálogo
    sin_fila: list = field(default_factory=list)          # productos del catálogo que esta hoja no cubre

    @property
    def titulo(self) -> str:
        fin = self.lunes + dt.timedelta(days=6)
        return f"{self.formato.titulo} — semana del {self.lunes:%d/%m} al {fin:%d/%m/%Y}"


def planes_de_la_semana(ins: pp.Insumos, lunes: dt.date, colchon: float, usar: bool) -> dict:
    """{fecha: Plan | None} de los 7 días (los planes diarios del sistema)."""
    salida = {}
    for i in range(7):
        d = lunes + dt.timedelta(days=i)
        r = pp.plan_del_dia(ins, d, colchon, usar)
        salida[d] = r[0] if r else None
    return salida


def _por_producto(planes: dict, lunes: dt.date) -> tuple[dict, dict]:
    """({clave: ([mat x 7], [vesp x 7])}, {clave: (opc1, opc2, opc3)}) a partir de los planes diarios."""
    turnos: dict = {}
    opciones: dict = {}
    for i in range(7):
        plan = planes.get(lunes + dt.timedelta(days=i))
        if plan is None:
            continue
        for f in plan.filas:
            k = plan_dia.clave_producto(f.producto)
            mat, vesp = turnos.setdefault(k, ([None] * 7, [None] * 7))
            mat[i] = sum(v for h, v in f.horas.items() if h <= 12)
            vesp[i] = sum(v for h, v in f.horas.items() if h >= 13)
            if any(t for t in f.tandas) and k not in opciones:
                opciones[k] = tuple(t or 0 for t in f.tandas)
    return turnos, opciones


def _suma(a, b):
    return a if b is None else (b if a is None else a + b)


def construir(formato: Formato, ins: pp.Insumos, planes: dict, lunes: dt.date, colchon: float, reserva: dict | None = None) -> PlanSemana:
    """El plan semanal de `formato` con las cifras de los planes diarios `planes` (ver planes_de_la_semana)."""
    claves = ins.catalogo.claves
    turnos, opciones = _por_producto(planes, lunes)
    faltan = [lunes + dt.timedelta(days=i) for i in range(7) if planes.get(lunes + dt.timedelta(days=i)) is None]
    filas: list = []
    sin_producto: list = []
    cubiertos: set = set()
    if formato.tipo == "decoracion":
        mapa = mapa_decoracion(formato, claves)
        grupos = {n: [k] if k else [] for n, k in mapa.items()}
    else:
        grupos = mapa_panaderia(formato, ins.catalogo.items)
    for tipo, nombre in formato.filas:
        if tipo == "grupo":
            filas.append(FilaSemana(nombre, grupo=True))
            continue
        fila = FilaSemana(nombre)
        miembros = grupos.get(nombre, [])
        if not miembros:
            sin_producto.append(nombre)
        for k in miembros:
            cubiertos.add(k)
            if k in turnos:
                for i in range(7):
                    fila.mat[i] = _suma(fila.mat[i], turnos[k][0][i])
                    fila.vesp[i] = _suma(fila.vesp[i], turnos[k][1][i])
            if k in opciones:
                fila.opciones = tuple(_suma(a, b) for a, b in zip(fila.opciones, opciones[k]))
        if reserva and reserva.get(nombre):
            fila.reserva = int(reserva[nombre])
        filas.append(fila)
    nombres = {i.clave: i.producto for i in ins.catalogo.items}
    return PlanSemana(formato=formato, lunes=lunes, filas=filas, colchon=colchon, faltan_dias=faltan, sin_producto=sin_producto,
                      sin_fila=[nombres[k] for k in claves if k not in cubiertos])


def totales(plan: PlanSemana) -> dict:
    items = [f for f in plan.filas if not f.grupo]
    suma = lambda col: sum(v or 0 for v in col)
    return {"mat": [suma([f.mat[i] for f in items]) for i in range(7)], "vesp": [suma([f.vesp[i] for f in items]) for i in range(7)],
            "opciones": tuple(suma([f.opciones[j] for f in items]) for j in range(3)), "reserva": suma([f.reserva for f in items])}


def _celda(v) -> str:
    return "" if not v else str(int(v))


def _pie(plan: PlanSemana) -> str:
    return (plan.formato.pie + f" · Plan hecho por el sistema solo con ventas (colchón {plan.colchon:.0%}) el {dt.date.today():%d/%m/%Y}.").strip(" ·")


# --------------------------------------------------------------------------------------------- salida: PDF y xlsx

_CSS = """
table { border-collapse: collapse; width: 786pt; font-family: sans-serif; font-size: 6.6pt; }
td, th { border: 0.4pt solid #000000; padding: 0.9pt 1.5pt; }
th { background-color: #375623; color: #ffffff; font-weight: bold; text-align: center; }
th.t { font-size: 9.5pt; padding: 3pt; }
td.n { text-align: center; } td.p { text-align: left; }
tr.g td { background-color: #c6e0b4; font-weight: bold; }
tr.alt td { background-color: #f2f8ee; }
tr.tot td { background-color: #375623; color: #ffffff; font-weight: bold; }
td.sep { border-left: 1.1pt solid #000000; }
p.pie { font-family: sans-serif; font-size: 6.2pt; margin-top: 3pt; }
"""


def _anchos(con_reserva: bool) -> list:
    if con_reserva:
        return [150] + [30] * 14 + [36] * 3 + [58]
    return [180] + [30] * 14 + [42] * 3


def _html(plan: PlanSemana, filas: list, ultima: bool) -> str:
    con_reserva = plan.formato.reserva is not None
    a = [x - 5 for x in _anchos(con_reserva)]
    n_col = 1 + 14 + 3 + (1 if con_reserva else 0)
    dias = "".join(f'<th colspan="2">{DIAS_CORTOS[i]} {(plan.lunes + dt.timedelta(days=i)):%d/%m}</th>' for i in range(7))
    cab = (f'<tr><th class="t" colspan="{n_col}">{plan.titulo}</th></tr>'
           f'<tr><th></th>{dias}<th colspan="3">{plan.formato.titulo_opciones}</th>' + ('<th></th>' if con_reserva else "") + '</tr>'
           f'<tr><th style="width:{a[0]}pt">{plan.formato.columna}</th>'
           + "".join(f'<th style="width:{a[1 + j]}pt">{"Mat" if j % 2 == 0 else "Vesp"}</th>' for j in range(14))
           + "".join(f'<th style="width:{a[15 + j]}pt">{e.replace(chr(10), " ")}</th>' for j, e in enumerate(plan.formato.etiquetas_opciones))
           + (f'<th style="width:{a[-1]}pt">{plan.formato.reserva.replace(chr(10), " ")}</th>' if con_reserva else "") + '</tr>')
    cuerpo = []
    visibles = 0
    for f in filas:
        if f.grupo:
            cuerpo.append(f'<tr class="g"><td colspan="{n_col}">{f.nombre}</td></tr>')
            visibles = 0
            continue
        celdas = "".join(f'<td class="n{" sep" if j % 2 == 0 else ""}">{_celda((f.mat if j % 2 == 0 else f.vesp)[j // 2])}</td>' for j in range(14))
        cuerpo.append(f'<tr class="{"alt" if visibles % 2 else ""}"><td class="p">{f.nombre}</td>{celdas}'
                      + "".join(f'<td class="n">{_celda(o)}</td>' for o in f.opciones)
                      + (f'<td class="n">{_celda(f.reserva)}</td>' if con_reserva else "") + "</tr>")
        visibles += 1
    if ultima:
        t = totales(plan)
        celdas = "".join(f'<td class="n{" sep" if j % 2 == 0 else ""}">{t["mat"][j // 2] if j % 2 == 0 else t["vesp"][j // 2]}</td>' for j in range(14))
        cuerpo.append('<tr class="tot"><td class="p">TOTAL</td>' + celdas + "".join(f'<td class="n">{x}</td>' for x in t["opciones"])
                      + (f'<td class="n">{t["reserva"]}</td>' if con_reserva else "") + "</tr>")
    pie = f'<p class="pie">{_pie(plan)}</p>' if ultima else ""
    return f"<table>{cab}{''.join(cuerpo)}</table>{pie}"


def plan_a_pdf(plan: PlanSemana, filas_por_pagina: int | None = None) -> bytes:
    """PDF horizontal A4 como la hoja del dueño (encabezado repetido, TOTAL y pie en la última página). Panadería cabe en una página;
    Decoración, en las que haga falta."""
    import pymupdf

    por_pagina = filas_por_pagina or (62 if plan.formato.tipo == "panaderia" else 46)
    paginas, actual = [], []
    for f in plan.filas:
        if len(actual) >= por_pagina and not f.grupo:
            huerfano = [actual.pop()] if actual and actual[-1].grupo else []      # un título de grupo no se queda solo al final de la página
            paginas.append(actual)
            actual = huerfano
        actual.append(f)
    paginas.append(actual)
    doc = pymupdf.open()
    for n, filas in enumerate(paginas):
        page = doc.new_page(width=842, height=595)
        page.insert_htmlbox(pymupdf.Rect(28, 20, 814, 578), _html(plan, filas, n == len(paginas) - 1), css=_CSS)
    return doc.tobytes()


def plan_a_xlsx(plan: PlanSemana) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    con_reserva = plan.formato.reserva is not None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = plan.formato.tipo.upper()[:31]
    verde = PatternFill("solid", fgColor=plan_dia.VERDE_OSCURO)
    claro = PatternFill("solid", fgColor=plan_dia.VERDE_CLARO)
    blanco = Font(bold=True, color="FFFFFF")
    borde = Border(*(Side(style="thin", color="000000"),) * 4)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n_col = 1 + 14 + 3 + (1 if con_reserva else 0)
    ws.cell(1, 1, plan.titulo)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_col)
    ws.cell(1, 1).font, ws.cell(1, 1).fill, ws.cell(1, 1).alignment = Font(bold=True, size=13, color="FFFFFF"), verde, centro
    ws.cell(2, 1, plan.formato.columna)
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    for i in range(7):
        c = 2 + 2 * i
        ws.cell(2, c, f"{DIAS_CORTOS[i]} {(plan.lunes + dt.timedelta(days=i)):%d/%m}")
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + 1)
        ws.cell(3, c, "Mat")
        ws.cell(3, c + 1, "Vesp")
    ws.cell(2, 16, plan.formato.titulo_opciones)
    ws.merge_cells(start_row=2, start_column=16, end_row=2, end_column=18)
    for j, e in enumerate(plan.formato.etiquetas_opciones):
        ws.cell(3, 16 + j, e)
    if con_reserva:
        ws.cell(2, 19, plan.formato.reserva)
        ws.merge_cells(start_row=2, start_column=19, end_row=3, end_column=19)
    for fila in ws.iter_rows(min_row=1, max_row=3, max_col=n_col):
        for c in fila:
            c.font, c.fill, c.alignment, c.border = blanco, verde, centro, borde
    r = 4
    visibles = 0
    for f in plan.filas:
        if f.grupo:
            ws.cell(r, 1, f.nombre)
            for c in ws[r][:n_col]:
                c.font, c.fill, c.border = Font(bold=True), PatternFill("solid", fgColor="C6E0B4"), borde
            visibles = 0
        else:
            ws.cell(r, 1, f.nombre).alignment = Alignment(horizontal="left")
            for i in range(7):
                ws.cell(r, 2 + 2 * i, f.mat[i] or None)
                ws.cell(r, 3 + 2 * i, f.vesp[i] or None)
            for j, o in enumerate(f.opciones):
                ws.cell(r, 16 + j, o or None)
            if con_reserva:
                ws.cell(r, 19, f.reserva or None)
            for c in ws[r][:n_col]:
                c.border = borde
                if c.column > 1:
                    c.alignment = Alignment(horizontal="center")
                if visibles % 2:
                    c.fill = claro
            visibles += 1
        r += 1
    t = totales(plan)
    ws.cell(r, 1, "TOTAL")
    for i in range(7):
        ws.cell(r, 2 + 2 * i, t["mat"][i])
        ws.cell(r, 3 + 2 * i, t["vesp"][i])
    for j, o in enumerate(t["opciones"]):
        ws.cell(r, 16 + j, o)
    if con_reserva:
        ws.cell(r, 19, t["reserva"])
    for c in ws[r][:n_col]:
        c.font, c.fill, c.border = blanco, verde, borde
        if c.column > 1:
            c.alignment = Alignment(horizontal="center")
    ws.cell(r + 2, 1, _pie(plan))
    ws.column_dimensions["A"].width = 30
    for col in range(2, n_col + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 7 if col < 19 else 11
    ws.freeze_panes = "B4"
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nombre_archivo(plan: PlanSemana, extension: str) -> str:
    tipo = "Panaderia" if plan.formato.tipo == "panaderia" else "Decoracion"
    return f"Plan_{tipo}_Semana_{plan.lunes.isoformat()}.{extension}"
